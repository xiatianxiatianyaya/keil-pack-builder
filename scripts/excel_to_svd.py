"""Normalized XLSX -> checked register model, Markdown description and CMSIS-SVD.
Requires openpyxl for read-only input. No hardware access or guessed register values.
"""
import argparse
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
import subprocess
import sys
import xml.etree.ElementTree as ET

COLS = {
    'Device': 'key value description',
    'Peripherals': 'name base_address description source',
    'Registers': 'peripheral name offset size access reset_value reset_mask modified_write_values read_action description source',
    'Fields': 'peripheral register name lsb width access modified_write_values read_action description source',
    'Enums': 'peripheral register field name value description source',
    'Interrupts': 'peripheral name value description source',
}
ACCESS = {'RO': 'read-only', 'WO': 'write-only', 'RW': 'read-write', 'W1': 'writeOnce', 'RW1': 'read-writeOnce'}
WRITES = {'oneToClear', 'oneToSet', 'oneToToggle', 'zeroToClear', 'zeroToSet', 'zeroToToggle', 'clear', 'set', 'modify'}
READS = {'clear', 'set', 'modify', 'modifyExternal'}
CPU_FLAGS = {'mpu_present': 'mpuPresent', 'fpu_present': 'fpuPresent', 'fpu_dp':'fpuDP',
             'dsp_present':'dspPresent', 'icache_present':'icachePresent', 'dcache_present':'dcachePresent',
             'itcm_present':'itcmPresent', 'dtcm_present':'dtcmPresent', 'vtor_present':'vtorPresent'}


def text(value):
    return '' if value is None else str(value).strip()


def integer(value, where):
    if isinstance(value, bool):
        raise ValueError(f'{where}: boolean is not an integer')
    if isinstance(value, (float, int)):
        if not isinstance(value, int) and (not value.is_integer()):
            raise ValueError(f'{where}: nonintegral number')
        if abs(value) > 2**53-1:
            raise ValueError(f'{where}: large Excel numeric cell; use exact hexadecimal text')
        n = int(value)
    elif re.fullmatch(r'0[xX][0-9a-fA-F]+|0[bB][01]+|[0-9]+', text(value)):
        s = text(value)
        n = int(s, 16 if s.lower().startswith('0x') else 2 if s.lower().startswith('0b') else 10)
    else:
        raise ValueError(f'{where}: expected nonnegative decimal/0xHEX/0bBIN, got {value!r}')
    if n < 0:
        raise ValueError(f'{where}: negative value')
    return n


def ident(value, where):
    s = text(value)
    if not re.fullmatch(r'[A-Za-z_][A-Za-z0-9_]*', s):
        raise ValueError(f'{where}: expected plain identifier; arrays/aliases need explicit adaptation')
    return s


def required(value, where):
    s = text(value)
    if not s:
        raise ValueError(f'{where}: required value missing')
    if any(ord(c) < 32 and c not in '\n\r\t' for c in s):
        raise ValueError(f'{where}: XML-invalid control character')
    return s


def boolean(value, where):
    s = text(value).lower()
    if s in ('true', '1'):
        return 'true'
    if s in ('false', '0'):
        return 'false'
    raise ValueError(f'{where}: expected true/false')


def access(value, where):
    s = text(value)
    s = ACCESS.get(s.upper(), s)
    if s not in ACCESS.values():
        raise ValueError(f'{where}: unknown access {value!r}; encode W1C as RW + oneToClear, not plain RW')
    return s


def effects(row, mode):
    where = row['_source']
    w, r = text(row['modified_write_values']), text(row['read_action'])
    if w and w not in WRITES or r and r not in READS:
        raise ValueError(f'{where}: unsupported write/read side effect')
    if w and mode == 'read-only' or r and mode in ('write-only', 'writeOnce'):
        raise ValueError(f'{where}: side effect conflicts with access')
    return w, r


def read_workbook(path):
    from openpyxl import load_workbook
    raw = Path(path).read_bytes()
    if Path(path).suffix.lower() != '.xlsx':
        raise ValueError('Use .xlsx; convert legacy .xls separately and preserve original')
    wb = load_workbook(BytesIO(raw), data_only=False, read_only=False, keep_links=False)
    rows = {}
    try:
        unknown = set(wb.sheetnames) - set(COLS) - {'Guide'}
        if unknown:
            raise ValueError(f'Unmapped sheets: {sorted(unknown)}; normalize a separate copy explicitly')
        for name, columns in COLS.items():
            if name not in wb.sheetnames:
                raise ValueError(f'Missing sheet {name}')
            ws = wb[name]
            if ws.merged_cells.ranges:
                raise ValueError(f'{name}: merged cells not supported; normalize first')
            data_rows = list(ws.iter_rows())
            if not data_rows:
                raise ValueError(f'{name}: empty sheet')
            # Only empty trailing cells may be ignored; nonempty unknown columns are rejected.
            header = [text(c.value) for c in data_rows[0]]
            while header and not header[-1]:
                header.pop()
            expected = columns.split()
            if len(header) != len(set(header)) or set(header) != set(expected):
                raise ValueError(f'{name}: columns must be exactly {expected}')
            rows[name] = []
            for line in data_rows[1:]:
                if all(c.value is None or c.value == '' for c in line):
                    continue
                for cell in line:
                    if cell.data_type in ('f', 'e'):
                        raise ValueError(f'{name}!{cell.coordinate}: formulas/errors not accepted; supply reviewed literal values')
                if any(c.value not in (None, '') for c in line[len(header):]):
                    raise ValueError(f'{name}!{line[0].row}: data outside named columns')
                row = {k: line[i].value for i, k in enumerate(header)}
                row['_source'] = f'{name}!{line[0].row}'
                rows[name].append(row)
    finally:
        wb.close()
    return rows, hashlib.sha256(raw).hexdigest()


def normalize(rows):
    meta, warnings = {}, []
    required_keys = {'name','version','description','address_unit_bits','width','cpu_name','revision','endian',
                     'mpu_present','fpu_present','nvic_prio_bits','vendor_systick_config'}
    allowed = required_keys | set(CPU_FLAGS)
    for row in rows['Device']:
        key = text(row['key'])
        if key not in allowed or key in meta:
            raise ValueError(f'{row["_source"]}: unknown/duplicate Device key {key}')
        meta[key] = required(row['value'], row['_source'])
    if required_keys - set(meta):
        raise ValueError(f'Missing Device keys: {sorted(required_keys-set(meta))}')
    meta['name'] = ident(meta['name'], 'Device.name')
    if integer(meta['address_unit_bits'], 'Device.address_unit_bits') != 8:
        raise ValueError('This converter supports byte-addressed devices only (address_unit_bits=8)')
    if integer(meta['width'], 'Device.width') not in (8,16,32,64):
        raise ValueError('Device.width must be 8/16/32/64')
    if not re.fullmatch(r'r\d{1,2}p\d{1,2}', meta['revision']):
        raise ValueError('Device.revision must be rNpM')
    if meta['endian'] not in ('little','big','selectable','other'):
        raise ValueError('Invalid endian')
    if not 1 <= integer(meta['nvic_prio_bits'], 'Device.nvic_prio_bits') <= 8:
        raise ValueError('nvic_prio_bits must be 1..8')
    for k in set(CPU_FLAGS) | {'vendor_systick_config'}:
        if k in meta:
            meta[k] = boolean(meta[k], f'Device.{k}')
    peripherals, registers, fields = {}, {}, {}
    for row in rows['Peripherals']:
        n = ident(row['name'], row['_source'])
        if n in peripherals:
            raise ValueError(f'{row["_source"]}: duplicate peripheral {n}')
        base = integer(row['base_address'], row['_source'])
        if base > 0xFFFFFFFF:
            raise ValueError('This converter supports 32-bit peripheral addresses only')
        peripherals[n] = dict(name=n, base_address=base, description=required(row['description'],row['_source']),
                              source=row['_source'], origin=text(row['source']), registers=[], interrupts=[])
    if not peripherals:
        raise ValueError('At least one peripheral required')
    address_ranges = []
    for row in rows['Registers']:
        pn, n = text(row['peripheral']), ident(row['name'],row['_source'])
        key = (pn,n)
        if pn not in peripherals or key in registers:
            raise ValueError(f'{row["_source"]}: orphan or duplicate register {key}')
        size = integer(row['size'],row['_source'])
        offset = integer(row['offset'],row['_source'])
        if size not in (8,16,32,64) or offset % (size//8):
            raise ValueError(f'{row["_source"]}: unsupported size or unaligned offset')
        mode = access(row['access'],row['_source'])
        rv, rm = integer(row['reset_value'],row['_source']), integer(row['reset_mask'],row['_source'])
        if rv >= 1<<size or rm >= 1<<size or rv & ~rm:
            raise ValueError(f'{row["_source"]}: reset overflow or set reset bits outside reset_mask')
        w,r = effects(row,mode)
        absolute = peripherals[pn]['base_address']+offset
        end = absolute+size//8
        if end > 1<<32 or absolute % (size//8):
            raise ValueError(f'{row["_source"]}: address overflow or misalignment')
        for start, stop, old in address_ranges:
            if absolute < stop and start < end:
                raise ValueError(f'{row["_source"]}: address overlap with {old}; aliases need explicit adaptation')
        address_ranges.append((absolute,end,key))
        reg = dict(name=n, offset=offset, address=absolute, size=size, access=mode,reset_value=rv,reset_mask=rm,
                   modified_write_values=w,read_action=r,description=required(row['description'],row['_source']),
                   source=row['_source'],origin=text(row['source']),fields=[])
        registers[key]=reg
        peripherals[pn]['registers'].append(reg)
    for row in rows['Fields']:
        key=(text(row['peripheral']),text(row['register']))
        name=ident(row['name'],row['_source'])
        fkey=key+(name,)
        if key not in registers or fkey in fields:
            raise ValueError(f'{row["_source"]}: orphan or duplicate field {fkey}')
        reg=registers[key]
        lsb,width=integer(row['lsb'],row['_source']),integer(row['width'],row['_source'])
        if not width or lsb+width>reg['size']:
            raise ValueError(f'{row["_source"]}: field out of register range')
        mask=((1<<width)-1)<<lsb
        if any(mask & f['mask'] for f in reg['fields']):
            raise ValueError(f'{row["_source"]}: overlapping fields')
        mode=access(row['access'],row['_source']) if text(row['access']) else reg['access']
        w,r=effects(row,mode)
        f=dict(name=name,lsb=lsb,width=width,mask=mask,access=mode,modified_write_values=w,read_action=r,
               description=required(row['description'],row['_source']),source=row['_source'],origin=text(row['source']),enums=[])
        fields[fkey]=f
        reg['fields'].append(f)
    for row in rows['Enums']:
        key=(text(row['peripheral']),text(row['register']),text(row['field']))
        if key not in fields:
            raise ValueError(f'{row["_source"]}: orphan enum {key}')
        f=fields[key]
        n,v=ident(row['name'],row['_source']),integer(row['value'],row['_source'])
        if v>=1<<f['width'] or any(e['name']==n or e['value']==v for e in f['enums']):
            raise ValueError(f'{row["_source"]}: enum overflow/duplicate; read/write-specific enums require adaptation')
        f['enums'].append(dict(name=n,value=v,description=required(row['description'],row['_source']),source=row['_source'],origin=text(row['source'])))
    irqs={}
    for row in rows['Interrupts']:
        pn=text(row['peripheral'])
        n,v=ident(row['name'],row['_source']),integer(row['value'],row['_source'])
        if pn not in peripherals or n in irqs or v in irqs.values() or v>479:
            raise ValueError(f'{row["_source"]}: orphan/duplicate/unsupported interrupt (shared IRQ needs one canonical row)')
        irqs[n]=v
        peripherals[pn]['interrupts'].append(dict(name=n,value=v,description=required(row['description'],row['_source']),source=row['_source'],origin=text(row['source'])))
    for p in peripherals.values():
        if not p['registers']:
            raise ValueError(f'{p["name"]}: peripheral has no registers')
        p['registers'].sort(key=lambda x:x['offset'])
        for r in p['registers']:
            r['fields'].sort(key=lambda x:x['lsb'])
            if not r['fields']:
                warnings.append(f'{p["name"]}.{r["name"]}: no fields supplied; whole register only')
    warnings.append('Flat register subset only; no arrays, clusters, derivedFrom, alternate registers, writeConstraint or read/write-specific enums. Adapt explicitly rather than flatten silently.')
    return dict(device=meta,peripherals=list(peripherals.values()),warnings=warnings)


def node(parent,tag,value):
    ET.SubElement(parent,tag).text=str(value)


def side_effects(parent,obj):
    for key,tag in [('modified_write_values','modifiedWriteValues'),('read_action','readAction')]:
        if obj[key]:
            node(parent,tag,obj[key])


def svd(model):
    d=model['device']
    root=ET.Element('device',schemaVersion='1.3')
    for key in ('name','version','description'):
        node(root,key,d[key])
    cpu=ET.SubElement(root,'cpu')
    for key,tag in [('cpu_name','name'),('revision','revision'),('endian','endian')]:
        node(cpu,tag,d[key])
    for key,tag in CPU_FLAGS.items():
        if key in d:
            node(cpu,tag,d[key])
    node(cpu,'nvicPrioBits',d['nvic_prio_bits'])
    node(cpu,'vendorSystickConfig',d['vendor_systick_config'])
    node(root,'addressUnitBits',d['address_unit_bits'])
    node(root,'width',d['width'])
    ps=ET.SubElement(root,'peripherals')
    for p in model['peripherals']:
        pn=ET.SubElement(ps,'peripheral')
        node(pn,'name',p['name']); node(pn,'description',p['description']); node(pn,'baseAddress',hex(p['base_address']))
        # Derived addressBlock covers supplied registers only, not claimed hardware aperture.
        ab=ET.SubElement(pn,'addressBlock')
        node(ab,'offset',0); node(ab,'size',hex(max(r['offset']+r['size']//8 for r in p['registers']))); node(ab,'usage','registers')
        for irq in p['interrupts']:
            ir=ET.SubElement(pn,'interrupt')
            for key in ('name','description','value'): node(ir,key,irq[key])
        rs=ET.SubElement(pn,'registers')
        for r in p['registers']:
            rn=ET.SubElement(rs,'register')
            for tag,val in [('name',r['name']),('description',r['description']),('addressOffset',hex(r['offset'])),
                            ('size',r['size']),('access',r['access']),('resetValue',hex(r['reset_value'])),('resetMask',hex(r['reset_mask']))]:
                node(rn,tag,val)
            side_effects(rn,r)
            if r['fields']:
                fs=ET.SubElement(rn,'fields')
                for f in r['fields']:
                    fn=ET.SubElement(fs,'field')
                    for tag,val in [('name',f['name']),('description',f['description']),('bitOffset',f['lsb']),('bitWidth',f['width']),('access',f['access'])]:
                        node(fn,tag,val)
                    side_effects(fn,f)
                    if f['enums']:
                        es=ET.SubElement(fn,'enumeratedValues')
                        for e in f['enums']:
                            en=ET.SubElement(es,'enumeratedValue')
                            for key in ('name','description','value'): node(en,key,e[key])
    ET.indent(root,space='  ')
    return ET.tostring(root,encoding='utf-8',xml_declaration=True)


def md_escape(value):
    return str(value).replace('&','&amp;').replace('<','&lt;').replace('>','&gt;').replace('|','&#124;').replace('\r','').replace('\n','<br>')


def describe(model):
    lines=[f'# {model["device"]["name"]} 寄存器描述','',md_escape(model['device']['description']),
           '', '由 Excel 标准化表生成；来源行保留于 model.json。addressBlock 是已提供寄存器范围的计算结果。未提供的位不自动定义为保留位。','']
    for p in model['peripherals']:
        lines += [f'## {p["name"]} — {hex(p["base_address"])}','',md_escape(p['description']),'']
        for r in p['registers']:
            lines += [f'### {r["name"]}','',md_escape(r['description']),'',
                      f'地址：`{hex(r["address"])}`；偏移：`{hex(r["offset"])}`；位宽：{r["size"]}；访问：{r["access"]}；复位：`{hex(r["reset_value"])}`；复位掩码：`{hex(r["reset_mask"])}`。',
                      f'寄存器写副作用：{r["modified_write_values"] or "未声明"}；读副作用：{r["read_action"] or "未声明"}。','',
                      '| 字段 | 位段 | 访问 | 写副作用 | 读副作用 | 描述 | 枚举 |','|---|---|---|---|---|---|---|']
            for f in r['fields']:
                vals=[f['name'],f'{f["lsb"]+f["width"]-1}:{f["lsb"]}',f['access'],f['modified_write_values'] or '继承/未声明',f['read_action'] or '继承/未声明',f['description'], '; '.join(f'{e["value"]}: {e["name"]} — {e["description"]}' for e in f['enums'])]
                lines.append('| '+' | '.join(md_escape(v) for v in vals)+' |')
            lines += ['']
        for irq in p['interrupts']:
            lines += [f'中断 `{irq["name"]}`：外部 IRQ 编号 {irq["value"]}，{md_escape(irq["description"])}。','']
    return '\n'.join(lines)+'\n'


def convert(path,out,svdconv=None):
    rows,digest=read_workbook(path)
    model=normalize(rows)
    model['source_workbook']={'filename':Path(path).name,'sha256':digest}
    outputs={f'{model["device"]["name"]}.svd':svd(model), 'registers.md':describe(model).encode('utf-8'),
             'model.json':json.dumps(model,ensure_ascii=False,indent=2).encode('utf-8')}
    report={'input_sha256':digest,'counts':{k:len(v) for k,v in rows.items()},'warnings':model['warnings'],
            'svdconv':{'status':'not_run'},'hardware':'not_run','errors':[],
            'created_at_utc':datetime.now(timezone.utc).isoformat()}
    out=Path(out)
    out.mkdir(parents=True,exist_ok=False)
    for name,data in outputs.items(): (out/name).write_bytes(data)
    if svdconv:
        try:
            exe=Path(svdconv).resolve(strict=True)
            cmd=[str(exe),str((out/f'{model["device"]["name"]}.svd').resolve()),'--strict']
            run=subprocess.run(cmd,cwd=out,capture_output=True,text=True,errors='replace',timeout=120)
            report['svdconv']={'status':'passed' if run.returncode==0 else 'review_required' if run.returncode==1 else 'failed',
                               'command':cmd,'tool_sha256':hashlib.sha256(exe.read_bytes()).hexdigest(),
                               'returncode':run.returncode,'stdout':run.stdout,'stderr':run.stderr}
            if run.returncode!=0: report['errors'].append('SVDConv did not return clean success; review log before integration')
        except (OSError,subprocess.SubprocessError) as e:
            report['svdconv']={'status':'failed','message':str(e)}
            report['errors'].append('SVDConv execution failed')
    report['output_sha256']={n:hashlib.sha256(data).hexdigest() for n,data in outputs.items()}
    (out/'conversion_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    return report


def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('xlsx'); ap.add_argument('--out',required=True); ap.add_argument('--svdconv')
    args=ap.parse_args()
    try:
        report=convert(args.xlsx,args.out,args.svdconv)
    except (ValueError,OSError,ImportError) as e:
        report={'errors':[f'{type(e).__name__}: {e}']}
    print(json.dumps(report,ensure_ascii=True,indent=2))
    return 1 if report['errors'] else 0


if __name__=='__main__':
    sys.exit(main())
